# Copyright 2026 the FyrQL authors
#
# This Source Code Form is subject to the terms of the Mozilla Public
# License, v. 2.0. If a copy of the MPL was not distributed with this
# file, You can obtain one at https://mozilla.org/MPL/2.0/.
#
# This Source Code Form is "Incompatible With Secondary Licenses", as
# defined by the Mozilla Public License, v. 2.0.

"""Writes workbooks with openpyxl and xlsxwriter into the directory given and prints, for each, what
read_xlsx reads from it: lines CASE <file> <arguments>, COLUMNS <name> <type> ..., ROW <value> ...,
with tabs between the parts, each value as the text a cast to VARCHAR gives it and escaped as Python's
unicode_escape does, NULL for a NULL. Prints SKIP if a writer is missing, and DONE at the end."""

import datetime as dt
import os
import sys

try:
    import openpyxl
    import xlsxwriter
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
except ImportError as e:
    print("SKIP " + str(e))
    sys.exit(0)

out = sys.argv[1]


def text(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float):
        return repr(v)
    return str(v).encode("unicode_escape").decode("ascii")


def case(name, args, columns, rows):
    print("CASE\t" + os.path.join(out, name) + "\t" + args)
    print("COLUMNS\t" + "\t".join(n + " " + t for n, t in columns))
    for row in rows:
        print("ROW\t" + "\t".join(text(v) for v in row))


# openpyxl writes inline strings, and dates as serial numbers with formats of their kind.
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "data"
ws.append(["n", "s", "b", "d", "ts", "t"])
rows = [
    [1.5, "plain", True, dt.date(2024, 2, 29), dt.datetime(2024, 2, 29, 13, 45, 30), dt.time(13, 45, 30)],
    [-2.0, "line1\nline2", False, dt.date(1900, 1, 15), dt.datetime(1900, 2, 28, 12, 0, 0), dt.time(0, 0, 1)],
    [None, "  spaces  ", None, dt.date(1999, 12, 31), dt.datetime(2100, 1, 1, 0, 0, 0), dt.time(23, 59, 59)],
    [1e300, "\u00e9t\u00e9 \U0001F600", True, dt.date(9999, 12, 31), dt.datetime(1970, 1, 1, 0, 0, 0), None],
]
for r in rows:
    ws.append(r)
wb.save(os.path.join(out, "openpyxl.xlsx"))
case("openpyxl.xlsx", "", [("n", "DOUBLE"), ("s", "VARCHAR"), ("b", "BOOLEAN"), ("d", "DATE"), ("ts", "TIMESTAMP"),
                           ("t", "TIME")],
     [[r[0], r[1], r[2], r[3], None if r[4] is None else r[4].strftime("%Y-%m-%d %H:%M:%S"), r[5]] for r in rows])

# Dates of 1904: the serial numbers count from 1904-01-01.
wb = openpyxl.Workbook()
wb.epoch = CALENDAR_MAC_1904
ws = wb.active
ws.append(["d", "ts"])
mac = [[dt.date(2024, 2, 29), dt.datetime(2024, 2, 29, 6, 30, 0)],
       [dt.date(1904, 1, 2), dt.datetime(1904, 1, 1, 0, 0, 1)]]
for r in mac:
    ws.append(r)
wb.save(os.path.join(out, "mac.xlsx"))
case("mac.xlsx", "", [("d", "DATE"), ("ts", "TIMESTAMP")],
     [[r[0], r[1].strftime("%Y-%m-%d %H:%M:%S")] for r in mac])

# xlsxwriter writes shared strings, escapes what XML cannot hold as _xHHHH_ and a literal _xHHHH_ as _x005F_xHHHH_.
wb = xlsxwriter.Workbook(os.path.join(out, "xlsxwriter.xlsx"))
ws = wb.add_worksheet("Sheet One")
date_format = wb.add_format({"num_format": "yyyy-mm-dd"})
stamp_format = wb.add_format({"num_format": "dd/mm/yyyy hh:mm:ss"})
ws.write_row(0, 0, ["s", "n", "b", "d", "ts", "f"])
strings = ["_x0041_ literal", "ctl\x01x", "tab\there", "", "z"]
for i, s in enumerate(strings, start=1):
    ws.write_string(i, 0, s)
    ws.write_number(i, 1, i * 0.25)
    ws.write_boolean(i, 2, i % 2 == 0)
    ws.write_datetime(i, 3, dt.datetime(2020, 1, i), date_format)
    ws.write_datetime(i, 4, dt.datetime(2020, 1, i, i, 30, 0), stamp_format)
    ws.write_formula(i, 5, "=B%d*4" % (i + 1), None, i)
wb.close()
case("xlsxwriter.xlsx", ", sheet = 'Sheet One'",
     [("s", "VARCHAR"), ("n", "DOUBLE"), ("b", "BOOLEAN"), ("d", "DATE"), ("ts", "TIMESTAMP"), ("f", "DOUBLE")],
     # xlsxwriter writes the empty string as a shared string without text, read as one.
     [[s, i * 0.25, i % 2 == 0, dt.date(2020, 1, i), "2020-01-0%d 0%d:30:00" % (i, i), float(i)]
      for i, s in enumerate(strings, start=1)])

# xlsxwriter in constant memory writes inline strings, row by row.
wb = xlsxwriter.Workbook(os.path.join(out, "streamed.xlsx"), {"constant_memory": True})
ws = wb.add_worksheet()
ws.write_row(0, 0, ["k", "v"])
for i in range(1, 1001):
    ws.write_number(i, 0, i)
    ws.write_string(i, 1, "v%d" % i)
wb.close()
case("streamed.xlsx", ", range = 'A1:B4'", [("k", "DOUBLE"), ("v", "VARCHAR")],
     [[float(i), "v%d" % i] for i in range(1, 4)])

print("DONE")
